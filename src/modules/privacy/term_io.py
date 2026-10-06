#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
term_io — Import/Export von Begriffslisten (Blacklist/Whitelist).

Unterstützt TXT (ein Begriff pro Zeile), CSV (erste Spalte) und XLSX
(erste Spalte des aktiven Blatts, via openpyxl). Rein lokal, ohne Qt.
"""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable, List

MAX_TERMS = 100_000


def _clean(terms: Iterable) -> List[str]:
    seen = set()
    result = []
    for term in terms:
        if term is None:
            continue
        text = str(term).strip()
        if not text or text.lower() == "nan" or text in seen:
            continue
        seen.add(text)
        result.append(text)
        if len(result) >= MAX_TERMS:
            break
    return result


def read_terms(path: str) -> List[str]:
    """Liest Begriffe aus einer Datei; wirft OSError/ValueError bei Fehlern."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()

    if suffix == ".xlsx":
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover - Abhängigkeit ist Pflicht
            raise ValueError("openpyxl ist nicht installiert") from exc
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            values = (row[0] for row in sheet.iter_rows(min_col=1, max_col=1, values_only=True))
            return _clean(values)
        finally:
            workbook.close()

    raw = file_path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - cp1252 dekodiert praktisch alles
        text = raw.decode("utf-8", errors="replace")

    if suffix == ".csv":
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = csv.reader(text.splitlines(), dialect)
        return _clean(row[0] for row in rows if row)

    return _clean(text.splitlines())


def write_terms(path: str, terms: Iterable[str]) -> None:
    """Schreibt Begriffe als TXT, CSV oder XLSX (nach Dateiendung)."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    items = sorted(_clean(terms), key=str.lower)

    if suffix == ".xlsx":
        import openpyxl

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Begriffe"
        for item in items:
            sheet.append([item])
        workbook.save(file_path)
        return

    if suffix == ".csv":
        with open(file_path, "w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            for item in items:
                writer.writerow([item])
        return

    file_path.write_text("\n".join(items) + ("\n" if items else ""), encoding="utf-8")
